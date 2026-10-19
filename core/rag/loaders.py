"""Load every document format into the shared Document schema, cleaned and ready to chunk.

One loader per format, each returning one Document per natural unit:
- PDF   one per page. Headings are found by font size; tables come out as Markdown tables (their
        text is not repeated as loose lines); repeated headers, footers and page numbers are dropped.
- DOCX  one per heading-2 section, numbered steps as "1. ", bullets as "- ", tables as Markdown.
        The page header ("Operations manual | Internal") is not body text and is skipped.
- PPTX  one per slide: title, bullets, any table, then the speaker notes.
- MD    one per "## " section; the text before the first one is the "Overview" section.

``clean_text`` runs on every unit: Unicode NFKC (PDF fonts extract "confirmed" with a one-character
"fi" ligature), hyphenated line breaks joined, non-breaking spaces and repeated spaces collapsed.
Document text is data: nothing here interprets it, and later steps treat it as untrusted.
"""
from __future__ import annotations

import re
import unicodedata
from collections import Counter
from pathlib import Path

from langchain_core.documents import Document

from core.docgen import DOCS_DIR, DocRecord, load_manifest
from core.rag.schema import make_document

_PAGE_NUMBER = re.compile(r"\bpage\s+\d+(\s+of\s+\d+)?\b", re.IGNORECASE)
_MARGIN_PT = 60  # text this close to the top or bottom edge, in a small font, is header or footer


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace(" ", " ").replace("­", "")      # non-breaking and soft hyphens
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)                   # "regis-\ntration" -> "registration"
    text = re.sub(r"[ \t]+", " ", text)
    text = "\n".join(line.strip() for line in text.splitlines())
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _table_markdown(rows: list[list[str]]) -> str:
    rows = [[" ".join((cell or "").split()).replace("|", "/") for cell in row] for row in rows if any(row)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    return "\n".join(lines + ["| " + " | ".join(r) + " |" for r in rows[1:]])


# --- PDF ------------------------------------------------------------------------------------------

def _pdf_blocks(page) -> list[dict]:
    """Text blocks and tables of one page, top to bottom: {y, size, text, table}."""
    found = page.find_tables().tables

    def inside(a, b):  # bbox a within bbox b
        return a != b and b[0] <= a[0] + 1 and b[1] <= a[1] + 1 and a[2] <= b[2] + 1 and a[3] <= b[3] + 1

    # The detector also reports a wrapped header cell as a tiny table of its own: keep real tables only.
    tables = [t for t in found if t.row_count > 1 and not any(inside(t.bbox, o.bbox) for o in found)]
    boxes = [t.bbox for t in found]
    items = [{"y": t.bbox[1], "size": 0.0, "text": _table_markdown(t.extract()), "table": True} for t in tables]
    for block in page.get_text("dict")["blocks"]:
        if block["type"] != 0:
            continue  # images
        x0, y0, x1, y1 = block["bbox"]
        if any(x0 < bx1 and bx0 < x1 and y0 < by1 and by0 < y1 for bx0, by0, bx1, by1 in boxes):
            continue  # inside a table, which is already in Markdown
        lines, size = [], 0.0
        for line in block["lines"]:
            lines.append("".join(span["text"] for span in line["spans"]))
            size = max([size] + [span["size"] for span in line["spans"]])
        text = ""
        for line in lines:  # a block is one paragraph: join its lines, mending hyphenated breaks
            text = text[:-1] + line if text.endswith("-") and line[:1].islower() else f"{text} {line}".strip()
        if text:
            items.append({"y": y0, "y1": y1, "size": size, "text": text, "table": False})
    return sorted(items, key=lambda item: item["y"])


def load_pdf(path: Path, record: DocRecord) -> list[Document]:
    import pymupdf

    with pymupdf.open(path) as pdf:
        pages = [(_pdf_blocks(page), page.rect.height) for page in pdf]

    # Body text is the size with the most characters; headings are clearly bigger.
    sizes = Counter()
    for blocks, _ in pages:
        for b in blocks:
            if not b["table"]:
                sizes[round(b["size"], 1)] += len(b["text"])
    body = sizes.most_common(1)[0][0] if sizes else 10.0

    def shape(text: str) -> str:  # page numbers vary, so compare lines with digits masked
        return re.sub(r"\d+", "#", text.lower())

    on_pages = Counter(shape(b["text"]) for blocks, _ in pages for b in {shape(x["text"]): x for x in blocks}.values())
    repeated = {s for s, n in on_pages.items() if len(pages) > 1 and n == len(pages)}

    documents, section = [], ""
    for number, (blocks, height) in enumerate(pages, start=1):
        parts, page_section = [], ""
        for b in blocks:
            if not b["table"]:
                in_margin = b["y"] < _MARGIN_PT or b.get("y1", 0) > height - _MARGIN_PT
                if shape(b["text"]) in repeated or (in_margin and b["size"] < body) or _PAGE_NUMBER.fullmatch(
                        b["text"].strip()):
                    continue  # header, footer or page number
                if b["size"] >= body * 1.6:
                    parts.append(f"# {b['text']}")
                    continue
                if b["size"] >= body * 1.15:
                    section = b["text"]
                    page_section = page_section or section
                    parts.append(f"## {b['text']}")
                    continue
            parts.append(b["text"])
        text = clean_text("\n\n".join(parts))
        if text:
            documents.append(make_document(text, record, page=number, section=page_section or section))
    return documents


# --- Word -----------------------------------------------------------------------------------------

def load_docx(path: Path, record: DocRecord) -> list[Document]:
    from docx import Document as WordDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    word = WordDocument(str(path))
    sections: list[tuple[str, list[str]]] = []
    step = 0
    for element in word.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            rows = [[cell.text for cell in row.cells] for row in Table(element, word).rows]
            if sections:
                sections[-1][1].append(_table_markdown(rows))
            continue
        if tag != "p":
            continue
        paragraph = Paragraph(element, word)
        text, style = paragraph.text.strip(), paragraph.style.name if paragraph.style is not None else ""
        if not text:
            continue
        if style == "Heading 1" or style == "Title":
            continue  # the document title is in the metadata and in every chunk's contextual header
        if style.startswith("Heading"):
            sections.append((text, [f"## {text}"]))
            step = 0
            continue
        if not sections:
            sections.append(("Overview", []))
        if style.startswith("List Number"):
            step += 1
            sections[-1][1].append(f"{step}. {text}")
        elif style.startswith("List"):
            sections[-1][1].append(f"- {text}")
        else:
            sections[-1][1].append(text)
    return [make_document(clean_text("\n".join(lines)), record, section=name)
            for name, lines in sections if len(lines) > 1 or not lines[0].startswith("## ")]


# --- PowerPoint -----------------------------------------------------------------------------------

def load_pptx(path: Path, record: DocRecord) -> list[Document]:
    from pptx import Presentation

    documents = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        lines = [f"## {title}"] if title else []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_table:
                lines.append(_table_markdown([[c.text for c in row.cells] for row in shape.table.rows]))
            elif shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        lines.append("  " * paragraph.level + f"- {text}")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Speaker notes: {notes}")
        text = clean_text("\n".join(lines))
        if text:
            documents.append(make_document(text, record, slide=number, section=title))
    return documents


# --- Markdown -------------------------------------------------------------------------------------

def load_markdown(path: Path, record: DocRecord) -> list[Document]:
    sections: list[tuple[str, list[str]]] = [("Overview", [])]
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):
            sections.append((line[3:].strip(), [line]))
        elif line.startswith("# ") and sections[0][1] == [] and len(sections) == 1:
            continue  # the file title; it is in the metadata and the contextual header
        else:
            sections[-1][1].append(line)
    return [make_document(text, record, section=name)
            for name, lines in sections if (text := clean_text("\n".join(lines)))]


# --- corpus ---------------------------------------------------------------------------------------

LOADERS = {"pdf": load_pdf, "docx": load_docx, "pptx": load_pptx, "md": load_markdown}


def load_document(record: DocRecord, docs_dir: Path) -> list[Document]:
    loader = LOADERS.get(record.source_type)
    if loader is None:
        raise ValueError(f"No loader for '{record.source_type}' ({record.path})")
    return loader(docs_dir / record.path, record)


def load_corpus(data_dir: str | Path) -> list[Document]:
    """Every document listed in ``<data_dir>/docs/manifest.json``, loaded and cleaned, in manifest order."""
    data_dir = Path(data_dir)
    documents: list[Document] = []
    for record in load_manifest(data_dir):
        documents.extend(load_document(record, data_dir / DOCS_DIR))
    return documents

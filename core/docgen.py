"""Generate the synthetic document corpus for RAG, consistent with the CSV dataset.

Five corpora, written under ``<data dir>/docs/``:
- event_reports/      one post-event report PDF per event (narrative, a sponsor results table,
                      repeated page headers and footers that ingestion must clean away)
- pricing/            a sponsorship pricing brochure PDF per year, with package and add-on tables
- sponsor_decks/      a partnership review PPTX, with speaker notes, for the six largest sponsors
- operations/         four operations SOPs as Word documents
- session_abstracts/  Markdown abstracts for the ten best-attended sessions of each event

Every figure comes from the CSVs through ``core.analysis``, so a document never contradicts a
dashboard. Qualitative facts (venue issues, lessons, sponsor requests, prices, SOP rules) exist
only in the documents, which is what makes retrieval necessary. ``eval/golden_set.json`` holds
20 questions built from those facts, with reference answers and the documents that answer them.

Randomness uses its own seeded stream, so the CSVs are unchanged and the same seed always
produces the same documents.
"""
from __future__ import annotations

import html
import io
import json
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

from core import analysis

CORPORA = ("event_reports", "pricing", "sponsor_decks", "operations", "session_abstracts")
DOCS_DIR, EVAL_DIR = "docs", "eval"
MANIFEST, GOLDEN_SET = "manifest.json", "golden_set.json"
ORGANISER = "ConferAtlas Events"
_FIXED_DATE = datetime(2026, 10, 1)  # document metadata date, so files are reproducible

VENUES = {
    "Austin": "Lone Star Convention Center", "Boston": "Harborfront Exhibition Hall",
    "Chicago": "Lakeshore Conference Center", "New York": "Empire Pier Conference Hall",
    "San Francisco": "Bayview Convention Center", "Toronto": "Lakeside Congress Centre",
}
# (issue, how it was handled), by format. Hybrid events draw from all three lists.
_ISSUES = {
    "in-person": (
        ("Wi-Fi in the main hall dropped for about 20 minutes during the opening keynote",
         "the venue switched the hall to a backup network and added two access points for day 2"),
        ("the registration queue peaked at 25 minutes on the first morning",
         "two extra badge printers were opened and a walk-in lane was added"),
        ("catering ran out of vegetarian lunches on day 1",
         "the caterer increased vegetarian portions by 40% for the remaining days"),
        ("the air conditioning in Hall B failed on the first afternoon",
         "sessions in Hall B moved to the ballroom until the system was repaired overnight"),
        ("signage to the level 2 breakout rooms confused attendees",
         "volunteers were posted at the escalators and new floor decals were added"),
    ),
    "virtual": (
        ("the streaming platform buffered for about 10 minutes during the opening keynote",
         "production switched to the backup encoder and the keynote recording was published the same day"),
        ("some attendees on corporate networks could not join breakout rooms",
         "the help desk shared a browser-only link that bypassed the blocked app"),
        ("the chat moderation queue fell behind during the live Q&A",
         "two more moderators were added and questions were collected in advance"),
        ("the virtual expo hall loaded slowly on mobile devices",
         "the platform vendor shipped a lighter mobile layout before day 2"),
    ),
    "hybrid": (
        ("remote attendees could not hear audience questions in two sessions",
         "handheld microphones were made mandatory for audience questions"),
    ),
}
_HIGHLIGHTS = (
    "the opening keynote drew the largest audience of the event",
    "the sponsor networking reception was the best-rated social event",
    "the new mobile app check-in cut queue times in half",
    "a live product demo theater was added and stayed full all day",
    "speed-mentoring sessions were fully booked within an hour",
)
# (lesson, recommendation for next time)
_LESSONS = (
    ("too many confirmed registrants did not show up",
     "send a reminder email 48 hours before the event, with the agenda and a calendar invite"),
    ("booth traffic dropped sharply after lunch",
     "move the expo hall next to the coffee stations so attendees pass the booths between sessions"),
    ("popular breakout rooms were overcrowded",
     "cap breakout room registration at 80% of capacity and stream overflow sessions"),
    ("technical problems were found only on the first morning",
     "run a full technical rehearsal the day before the event"),
    ("attendees asked repeatedly for the session slides",
     "publish session slides within 24 hours of the event"),
)
_SPONSOR_REQUESTS = (
    "asked for a larger booth near the coffee station",
    "asked to move their demo slot to the morning",
    "asked for attendee job titles in the lead export",
    "asked for a speaking slot in the main track next year",
    "asked for more signage pointing to their booth",
)

# Pricing: doc-only facts. 2026 prices are 2025 prices plus 6%, rounded to $500 (to $50 below $5,000).
_PACKAGE_INFO = (  # (package, format, price 2025 USD, booth, includes)
    ("Title Sponsor", "in-person / hybrid", 95_000, "20 x 20 ft island",
     "Keynote introduction, logo on all signage and the event app, 10 full passes, lead retrieval for 4 staff"),
    ("Premier Booth", "in-person / hybrid", 48_000, "10 x 20 ft",
     "Logo in the event app, 6 full passes, lead retrieval for 3 staff, one demo theater slot"),
    ("Standard Booth", "in-person / hybrid", 22_000, "10 x 10 ft",
     "Listing in the event app, 3 full passes, lead retrieval for 2 staff"),
    ("Exhibitor Table", "in-person / hybrid", 9_500, "6 ft table",
     "Listing in the event app, 2 full passes, lead retrieval for 1 staff"),
    ("Digital Premier", "virtual", 30_000, "Premium virtual booth",
     "Banner on the virtual lobby, a sponsored breakout session, attendee chat and lead export"),
    ("Virtual Booth", "virtual", 12_000, "Standard virtual booth",
     "Listing in the virtual expo hall, attendee chat and lead export"),
)
_ADD_ONS = (("Keynote speaking slot (20 minutes)", 15_000), ("Sponsored breakout session", 8_000),
            ("Branded lanyards", 6_000), ("Extra lead retrieval licence", 750))

_TITLE_PATTERNS = (
    "{track}: lessons from the field", "The future of {low}", "A practical {low} playbook",
    "Scaling {low} without burning out your team", "{track} myths, busted",
    "Case study: {low} at a mid-size organization", "Measuring what matters in {low}",
    "What we got wrong about {low}",
)
_FOCUS = ("a step-by-step framework drawn from recent projects", "results from a year-long pilot",
          "the three mistakes teams make most often", "a live walkthrough of real examples",
          "benchmarks from more than fifty organizations")
_TAKEAWAY = ("a checklist they can use the following week", "templates to adapt for their own team",
             "a clear way to explain the business case", "practical metrics to track progress",
             "a reading list and next steps")


@dataclass
class DocRecord:
    doc_id: str
    corpus: str
    source_type: str  # pdf, pptx, docx or md
    path: str         # relative to the docs folder, with forward slashes
    title: str
    event_id: str = ""
    event_type: str = ""
    city: str = ""
    year: int | None = None
    sponsor_id: str = ""


@dataclass
class _EventFacts:
    issue: str
    resolution: str
    highlight: str
    lesson: str
    recommendation: str
    request_sponsor: str
    request: str
    venue: str


@dataclass
class _Corpus:
    records: list[DocRecord] = field(default_factory=list)
    event_facts: dict[str, _EventFacts] = field(default_factory=dict)
    proposals: dict[str, str] = field(default_factory=dict)      # sponsor_id -> next-year proposal
    abstracts: dict[str, list[dict]] = field(default_factory=dict)  # event_id -> sessions written


# --- helpers ------------------------------------------------------------------------------

def _money(value: float) -> str:
    return f"${value:,.0f}"


def _pct(value: float | None) -> str:
    return "n/a" if value is None or pd.isna(value) else f"{value:.0%}"


def _year_price(price_2025: int, year: int) -> int:
    if year <= 2025:
        return price_2025
    step = 500 if price_2025 >= 5_000 else 50
    return int(round(price_2025 * 1.06 / step) * step)


def _dates(row) -> str:
    start, end = pd.Timestamp(row.start_date), pd.Timestamp(row.end_date)
    if start == end:
        return start.strftime("%B %d, %Y").replace(" 0", " ")
    return f"{start.strftime('%B %d').replace(' 0', ' ')} to {end.strftime('%B %d, %Y').replace(' 0', ' ')}"


def _lower(name: str) -> str:
    """Lower-case a track name for use mid-sentence, keeping acronyms and brand-style words:
    "AI & ML" and "DevOps" stay, "Cloud" -> "cloud"."""
    return " ".join(w if any(c.isupper() for c in w[1:]) or w.isupper() else w.lower() for w in name.split())


def _pick(rng: np.random.Generator, options):
    return options[int(rng.integers(len(options)))]


def _write_pdf(path: Path, body_html: str, header: str, footer: str, title: str) -> None:
    """Render HTML to a Letter-size PDF, then stamp a header and a numbered footer on every page."""
    import pymupdf

    css = ("body {font-family: sans-serif; font-size: 10pt; line-height: 1.35;} h1 {font-size: 18pt;} "
           "h2 {font-size: 13pt; margin-top: 12pt;} table {border-collapse: collapse; width: 100%;} "
           "th, td {border: 1px solid #999; padding: 3pt; font-size: 9pt;} th {background-color: #e8eef6;}")
    story = pymupdf.Story(html=body_html, user_css=css)
    buffer = io.BytesIO()  # rendered in memory: Windows locks a temporary file while it is open
    writer = pymupdf.DocumentWriter(buffer)
    page_rect = pymupdf.paper_rect("letter")
    content = page_rect + (54, 72, -54, -72)
    more = True
    while more:
        device = writer.begin_page(page_rect)
        more, _ = story.place(content)
        story.draw(device)
        writer.end_page()
    writer.close()

    doc = pymupdf.open("pdf", buffer.getvalue())
    for number, page in enumerate(doc, start=1):
        page.insert_text((54, 40), header, fontsize=8, color=(0.4, 0.4, 0.4))
        page.insert_text((54, page_rect.height - 36), f"{footer} | Page {number} of {doc.page_count}",
                         fontsize=8, color=(0.4, 0.4, 0.4))
    stamp = "D:" + _FIXED_DATE.strftime("%Y%m%d%H%M%S")
    doc.set_metadata({"title": title, "author": ORGANISER, "creationDate": stamp, "modDate": stamp,
                      "producer": "ConferAtlas generator", "creator": "ConferAtlas generator"})
    doc.save(str(path), garbage=3, deflate=True)
    doc.close()


# --- event reports ----------------------------------------------------------------------------

def _event_reports(rng, tables, out: Path, corpus: _Corpus) -> None:
    kpis = analysis.event_kpis(tables).set_index("event_id")
    sessions = analysis.session_performance(tables)
    results = analysis.sponsorship_results(tables)
    regs = tables["registrations.csv"]
    lessons = list(rng.permutation(len(_LESSONS)))
    highlights = list(rng.permutation(len(_HIGHLIGHTS)))
    used_issues: set[str] = set()

    for i, ev in enumerate(tables["events.csv"].sort_values("start_date").itertuples()):
        k = kpis.loc[ev.event_id]
        pool = [*_ISSUES[ev.format]] if ev.format != "hybrid" else [*_ISSUES["hybrid"], *_ISSUES["in-person"],
                                                                     *_ISSUES["virtual"]]
        fresh = [p for p in pool if p[0] not in used_issues] or pool
        issue, resolution = _pick(rng, fresh)
        used_issues.add(issue)
        lesson, recommendation = _LESSONS[lessons[i % len(_LESSONS)]]
        ev_results = results[results["event_id"] == ev.event_id].sort_values("leads", ascending=False)
        requester = ev_results.iloc[int(rng.integers(len(ev_results)))]
        request = _SPONSOR_REQUESTS[i % len(_SPONSOR_REQUESTS)]
        venue = (f"the ConferAtlas Live virtual platform, produced from our {ev.city} studio"
                 if ev.format == "virtual" else VENUES[ev.city])
        corpus.event_facts[ev.event_id] = _EventFacts(issue, resolution, _HIGHLIGHTS[highlights[i % 5]], lesson,
                                                      recommendation, requester["sponsor"], request, venue)
        facts = corpus.event_facts[ev.event_id]

        ev_sessions = sessions[(sessions["event_id"] == ev.event_id) & (sessions["attendance"] >= 20)]
        best = ev_sessions.sort_values(["rating", "attendance"], ascending=False).head(3)
        worst = ev_sessions.sort_values(["rating", "attendance"]).iloc[0]
        ev_regs = regs[(regs["event_id"] == ev.event_id) & regs["attended"].astype(bool)]
        roles = ev_regs["attendee_role"].value_counts().head(3)
        countries = ev_regs["country"].value_counts(normalize=True).head(3)
        variance = analysis._div(ev.actual_cost - ev.budget, ev.budget)
        e = html.escape

        rows = "".join(
            f"<tr><td>{e(r.sponsor)}</td><td>{e(r.package)}</td><td>{r.leads}</td><td>{r.leads_target}</td>"
            f"<td>{_pct(r.attainment)}</td></tr>" for r in ev_results.itertuples())
        body = f"""
<h1>{e(ev.name)}</h1>
<p><b>Post-event report</b> | {_dates(ev)} | {e(ev.city)} | {ev.format} {ev.event_type} event | Venue: {e(facts.venue)}</p>
<h2>1. Executive summary</h2>
<p>The {e(ev.name)} welcomed {k.attended:,} attendees out of {k.confirmed:,} confirmed registrations
({k.registrations:,} registrations in total), a show-up rate of {_pct(k.show_up_rate)}. Attendees rated the
program {k.avg_rating:.2f} out of 5 on average, for a net promoter score of {k.nps:.0f}. {k.sponsors} sponsors
took part; {k.sponsors_met_target} of them met their lead target, and sponsors collected {k.leads:,} leads
against a combined target of {k.leads_target:,}. The event cost {_money(ev.actual_cost)} against a budget of
{_money(ev.budget)} ({variance:+.1%}).</p>
<h2>2. Audience</h2>
<p>The largest attendee groups were {", ".join(f"{e(r)} ({n:,})" for r, n in roles.items())}. Most attendees
came from {", ".join(f"{e(c)} ({v:.0%})" for c, v in countries.items())}.</p>
<h2>3. Program highlights</h2>
<p>The program ran {k.sessions} sessions. The highest-rated sessions were
{"; ".join(f"{e(s.track)} with {e(s.speaker)} (rated {s.rating:.2f}, {s.attendance} attendees)" for s in best.itertuples())}.
The lowest-rated session was {e(worst.track)} with {e(worst.speaker)} (rated {worst.rating:.2f}).
A highlight of the event: {facts.highlight}.</p>
<h2>4. Venue and operations</h2>
<p>The main operational issue was that {facts.issue}. In response, {facts.resolution}.</p>
<h2>5. Sponsor results</h2>
<table><tr><th>Sponsor</th><th>Package</th><th>Leads</th><th>Target</th><th>Attainment</th></tr>{rows}</table>
<p>Sponsor feedback: {e(facts.request_sponsor)} {facts.request}.</p>
<h2>6. Lessons learned</h2>
<p>Lesson: {facts.lesson}. Recommendation for the next edition: {facts.recommendation}.</p>
"""
        rel = f"event_reports/{ev.event_id}_post_event_report.pdf"
        title = f"{ev.name}: post-event report"
        _write_pdf(out / rel, body, header=f"{ORGANISER} | {ev.name}",
                   footer=f"{ORGANISER} | Post-event report | Confidential", title=title)
        corpus.records.append(DocRecord(f"report_{ev.event_id}", "event_reports", "pdf", rel, title, ev.event_id,
                                        ev.event_type, ev.city, pd.Timestamp(ev.start_date).year))


# --- pricing brochures ------------------------------------------------------------------------

def _pricing(out: Path, corpus: _Corpus) -> None:
    targets = {"Title Sponsor": 220, "Premier Booth": 140, "Standard Booth": 80, "Exhibitor Table": 40,
               "Digital Premier": 160, "Virtual Booth": 70}  # per 1,000 attendees, as in core.generator
    for year in (2025, 2026):
        rows = "".join(
            f"<tr><td>{p}</td><td>{fmt}</td><td>{_money(_year_price(price, year))}</td><td>{booth}</td>"
            f"<td>{targets[p]}</td><td>{inc}</td></tr>" for p, fmt, price, booth, inc in _PACKAGE_INFO)
        add_ons = "".join(f"<tr><td>{name}</td><td>{_money(_year_price(price, year))}</td></tr>"
                          for name, price in _ADD_ONS)
        guarantee = ("<h2>Lead guarantee (new for 2026)</h2><p>Title Sponsor and Premier Booth sponsors whose leads "
                     "reach less than 70% of their lead target receive a credit of 15% of their package price "
                     "toward the following year.</p>") if year >= 2026 else ""
        body = f"""
<h1>Sponsorship packages {year}</h1>
<p>{ORGANISER} runs technology, business, medical and education conferences in person, online and in hybrid
format. Sponsors reach decision makers through booths, sessions and digital placements, and receive a lead
export after every event.</p>
<h2>Packages</h2>
<table><tr><th>Package</th><th>Format</th><th>Price</th><th>Booth</th><th>Lead target per 1,000 attendees</th>
<th>Includes</th></tr>{rows}</table>
<h2>Add-ons</h2>
<table><tr><th>Add-on</th><th>Price</th></tr>{add_ons}</table>
{guarantee}
<h2>Payment and cancellation terms</h2>
<p>A 50% deposit is due at signing and the balance 60 days before the event. Cancellations more than 90 days
before the event receive a 75% refund of fees paid; cancellations 30 to 90 days before receive 25%; cancellations
within 30 days are not refunded. Lead targets are scaled to each event's expected attendance and confirmed in the
sponsorship agreement.</p>
"""
        rel = f"pricing/sponsorship_packages_{year}.pdf"
        title = f"Sponsorship packages {year}"
        _write_pdf(out / rel, body, header=f"{ORGANISER} | Sponsorship packages {year}",
                   footer=f"{ORGANISER} | Pricing valid for events held in {year}", title=title)
        corpus.records.append(DocRecord(f"pricing_{year}", "pricing", "pdf", rel, title, year=year))


# --- sponsor decks ----------------------------------------------------------------------------

def _sponsor_decks(tables, out: Path, corpus: _Corpus) -> None:
    from pptx import Presentation

    cards = analysis.sponsor_scorecards(tables)
    results = analysis.sponsorship_results(tables)
    leads = tables["leads.csv"]
    top = cards.sort_values(["contract_value", "sponsor"], ascending=[False, True]).head(6)
    for card in top.itertuples():
        mine = results[results["sponsor_id"] == card.sponsor_id].sort_values("start_date")
        if card.attainment < analysis.LEAD_TARGET_ALERT:
            proposal = ("add a pre-event email to attendees and a demo theater slot at every event, "
                        "to lift leads back above target")
        elif card.attainment >= 1.1:
            proposal = (f"upgrade to the next package tier at {card.best_event}'s next edition, "
                        "and add a keynote speaking slot")
        else:
            proposal = "keep the current packages and add a sponsored breakout session at each event"
        corpus.proposals[card.sponsor_id] = proposal
        year = int(pd.Timestamp(mine["start_date"].max()).year)

        deck = Presentation()
        deck.core_properties.title = f"{card.sponsor}: partnership review"
        deck.core_properties.author = ORGANISER
        deck.core_properties.created = deck.core_properties.modified = _FIXED_DATE
        title_layout, bullets_layout = deck.slide_layouts[0], deck.slide_layouts[1]

        def slide(title: str, lines: list[str], notes: str, layout=bullets_layout):
            s = deck.slides.add_slide(layout)
            s.shapes.title.text = title
            body = s.placeholders[1].text_frame
            body.text = lines[0]
            for line in lines[1:]:
                body.add_paragraph().text = line
            s.notes_slide.notes_text_frame.text = notes

        slide(f"{card.sponsor} x {ORGANISER}", [f"Partnership review, events through {year}"],
              f"Opening slide for the {card.sponsor} renewal meeting.", title_layout)
        slide(f"About {card.sponsor}", [f"Industry: {card.industry}", f"Tier: {card.tier}",
                                        f"Goals: {card.goals}", f"Contract value: {_money(card.contract_value)}"],
              f"{card.sponsor} told us their main goals are {card.goals.lower()}.")
        slide("Events sponsored", [f"{r.event}: {r.package}, {r.leads} leads vs target {r.leads_target}"
                                   for r in mine.itertuples()],
              "Lead counts come from the badge-scan exports for each event.")
        slide("Results", [f"Total leads: {card.leads:,} against a target of {card.leads_target:,} "
                          f"({_pct(card.attainment)})",
                          f"Cost per lead: {_money(card.cost_per_lead)}",
                          f"Best event: {card.best_event} ({card.best_event_leads} leads)",
                          f"Targets met at {card.targets_met} of {card.events} events"],
              "Cost per lead is the contract value divided by total leads across all events.")
        notes = analysis.lead_notes(leads[leads["sponsor_id"] == card.sponsor_id], n=3)
        slide("What leads told us", [f"{n.note} ({n.leads} leads)" for n in notes.itertuples()],
              "Most common notes that booth staff recorded when scanning leads.")
        slide("Proposal for next year", [proposal[0].upper() + proposal[1:]],
              f"Our proposal to {card.sponsor}: {proposal}. Present this as the main ask of the meeting.")

        rel = f"sponsor_decks/{card.sponsor_id}_partnership_review.pptx"
        deck.save(str(out / rel))
        corpus.records.append(DocRecord(f"deck_{card.sponsor_id}", "sponsor_decks", "pptx", rel,
                                        f"{card.sponsor}: partnership review", year=year,
                                        sponsor_id=card.sponsor_id))


# --- operations SOPs ----------------------------------------------------------------------------

_SOPS = {
    "venue_setup": ("SOP-01 Venue setup", {
        "Purpose": ["This procedure covers preparing an in-person or hybrid venue before doors open."],
        "Load-in": ["Load-in starts at 6:00 AM on the day before the event.",
                    "The audio-visual team completes sound and projector checks in every room by 4:00 PM.",
                    "Capacity signs approved by the fire marshal are posted at every room entrance."],
        "Network": ["Test the Wi-Fi with at least 200 concurrent devices in the main hall.",
                    "Keep a backup network ready for the main hall and keynote stage."],
        "Sign-off": ["The venue lead signs the readiness checklist by 6:00 PM and sends it to the event director."],
    }),
    "registration_desk": ("SOP-02 Registration desk", {
        "Purpose": ["This procedure covers badge pickup, walk-ins and queue management."],
        "Opening": ["The registration desk opens 90 minutes before the first session.",
                    "At least one badge printer per 300 expected attendees must be running."],
        "Queue management": ["The target wait time is under 10 minutes.",
                             "If the queue exceeds 15 minutes, open the overflow lane and call the floor manager."],
        "Special cases": ["Walk-in registrations pay at the desk by card only.",
                          "A lost badge is reprinted only after the attendee shows photo ID.",
                          "Speakers and VIPs use the dedicated lane at the left of the desk."],
    }),
    "booth_logistics": ("SOP-03 Sponsor booth logistics", {
        "Purpose": ["This procedure covers sponsor move-in, lead retrieval and teardown."],
        "Move-in": ["Booth move-in runs from 2:00 PM to 8:00 PM on the day before the event.",
                    "Freight must be labelled with the sponsor name and booth number."],
        "Lead retrieval": ["Badge scanners must be charged and tested by 7:00 AM on each event day.",
                           "Each sponsor receives a lead export within 48 hours of the event closing.",
                           "The export is a CSV with name, company, job title, source and notes."],
        "Teardown": ["Teardown may not start before 4:00 PM on the last day.",
                     "All booths must be cleared by 8:00 PM; late teardown is billed at $500 per hour."],
    }),
    "virtual_streaming": ("SOP-04 Virtual streaming", {
        "Purpose": ["This procedure covers streaming for virtual and hybrid events."],
        "Rehearsal": ["Run a full technical rehearsal with every speaker the day before the event.",
                      "A backup encoder must be running in parallel for the keynote stream."],
        "During the event": ["The target stream latency is under 15 seconds.",
                             "If a stream is down for more than 5 minutes, the producer switches to the backup "
                             "encoder and escalates to the platform vendor's on-call engineer."],
        "After the event": ["Session recordings are published within 24 hours.",
                            "Chat logs are deleted after 30 days."],
    }),
}


def _operations(out: Path, corpus: _Corpus) -> None:
    from docx import Document

    for key, (title, sections) in _SOPS.items():
        doc = Document()
        doc.core_properties.title = title
        doc.core_properties.author = ORGANISER
        doc.core_properties.created = doc.core_properties.modified = _FIXED_DATE
        doc.sections[0].header.paragraphs[0].text = f"{ORGANISER} | Operations manual | Internal"
        doc.add_heading(title, level=1)
        for heading, lines in sections.items():
            doc.add_heading(heading, level=2)
            style = "Normal" if heading == "Purpose" else "List Number"
            for line in lines:
                doc.add_paragraph(line, style=style)
        rel = f"operations/{key}.docx"
        doc.save(str(out / rel))
        corpus.records.append(DocRecord(f"sop_{key}", "operations", "docx", rel, title))


# --- session abstracts ----------------------------------------------------------------------------

def _session_abstracts(rng, tables, out: Path, corpus: _Corpus) -> None:
    sessions = tables["sessions.csv"]
    for ev in tables["events.csv"].sort_values("start_date").itertuples():
        top = sessions[sessions["event_id"] == ev.event_id].sort_values(["attendance", "session_id"],
                                                                        ascending=[False, True]).head(10)
        lines, written, used = [f"# {ev.name}: session abstracts", "",
                                f"{_dates(ev)}, {ev.city}. The ten best-attended sessions.", ""], [], set()
        for s in top.itertuples():
            for _ in range(len(_TITLE_PATTERNS)):
                title = _pick(rng, _TITLE_PATTERNS).format(track=s.track, low=_lower(s.track))
                if title not in used:
                    break
            used.add(title)
            lines += [f"## {title}", "", f"- **Track:** {s.track}", f"- **Speaker:** {s.speaker}",
                      f"- **Format:** {s.mode}", "",
                      f"In this {s.mode} session, {s.speaker} shares {_pick(rng, _FOCUS)} on {_lower(s.track)}. "
                      f"Attendees will leave with {_pick(rng, _TAKEAWAY)}.", ""]
            written.append({"session_id": s.session_id, "title": title, "track": s.track, "speaker": s.speaker})
        corpus.abstracts[ev.event_id] = written
        rel = f"session_abstracts/{ev.event_id}_session_abstracts.md"
        (out / rel).write_text("\n".join(lines), encoding="utf-8", newline="\n")
        corpus.records.append(DocRecord(f"abstracts_{ev.event_id}", "session_abstracts", "md", rel,
                                        f"{ev.name}: session abstracts", ev.event_id, ev.event_type, ev.city,
                                        pd.Timestamp(ev.start_date).year))


# --- golden set -----------------------------------------------------------------------------------

def _golden_set(tables, corpus: _Corpus) -> list[dict]:
    events = tables["events.csv"].sort_values("start_date").reset_index(drop=True)
    name = dict(zip(events["event_id"], events["name"]))
    kpis = analysis.event_kpis(tables).set_index("event_id")
    results = analysis.sponsorship_results(tables)
    facts = corpus.event_facts
    ids = list(events["event_id"])
    decks = [r for r in corpus.records if r.corpus == "sponsor_decks"]
    sponsor_name = dict(zip(tables["sponsors.csv"]["sponsor_id"], tables["sponsors.csv"]["name"]))
    questions: list[dict] = []

    def add(question, category, corpora, docs, answer, answerable=True):
        questions.append({"id": f"Q{len(questions) + 1:02d}", "question": question, "category": category,
                          "answerable": answerable, "expected_corpora": corpora, "expected_doc_ids": docs,
                          "reference_answer": answer})

    a, b, c, d, x, y = ids[0], ids[3], ids[5], ids[7], ids[9], ids[11]
    add(f"What operational problem came up at the {name[a]}, and how was it handled?", "single_doc",
        ["event_reports"], [f"report_{a}"], f"{facts[a].issue[0].upper()}{facts[a].issue[1:]}; {facts[a].resolution}.")
    add(f"What lesson did the team record after the {name[b]}?", "single_doc", ["event_reports"],
        [f"report_{b}"], f"{facts[b].lesson[0].upper()}{facts[b].lesson[1:]}. "
                         f"Recommendation: {facts[b].recommendation}.")
    r = results[results["event_id"] == c].sort_values("leads", ascending=False).iloc[0]
    add(f"How many leads did {r.sponsor} collect at the {name[c]}, and what was their target?", "table_lookup",
        ["event_reports"], [f"report_{c}"], f"{r.leads} leads against a target of {r.leads_target}.")
    add(f"Which sponsor at the {name[d]} {facts[d].request}?", "single_doc", ["event_reports"],
        [f"report_{d}"], f"{facts[d].request_sponsor}.")
    premier_2026 = _year_price(48_000, 2026)
    add("How much does a Premier Booth cost for 2026 events?", "table_lookup", ["pricing"], ["pricing_2026"],
        f"{_money(premier_2026)}.")
    add("What does the Title Sponsor package include?", "table_lookup", ["pricing"], ["pricing_2026", "pricing_2025"],
        _PACKAGE_INFO[0][4] + ", with a 20 x 20 ft island booth.")
    add("If a sponsor cancels 60 days before the event, how much of the fees paid is refunded?", "single_doc",
        ["pricing"], ["pricing_2026", "pricing_2025"], "25% of fees paid (cancellations 30 to 90 days before).")
    add("What happens if a Premier Booth sponsor reaches less than 70% of its lead target?", "single_doc",
        ["pricing"], ["pricing_2026"], "From 2026 they receive a credit of 15% of the package price toward the "
                                       "following year (the lead guarantee).")
    add("By what time must badge scanners be charged and tested?", "single_doc", ["operations"],
        ["sop_booth_logistics"], "By 7:00 AM on each event day.")
    add("How soon after an event do sponsors receive their lead export, and what does it contain?", "single_doc",
        ["operations"], ["sop_booth_logistics"], "Within 48 hours of the event closing, as a CSV with name, "
                                                 "company, job title, source and notes.")
    add("What should the producer do if a stream is down for more than 5 minutes?", "single_doc", ["operations"],
        ["sop_virtual_streaming"], "Switch to the backup encoder and escalate to the platform vendor's on-call "
                                   "engineer.")
    deck = decks[0]
    add(f"What did we propose to {sponsor_name[deck.sponsor_id]} for next year?", "single_doc", ["sponsor_decks"],
        [deck.doc_id], corpus.proposals[deck.sponsor_id][0].upper() + corpus.proposals[deck.sponsor_id][1:] + ".")
    s = corpus.abstracts[x][0]
    add(f"Who presented the session \"{s['title']}\" at the {name[x]}?", "single_doc", ["session_abstracts"],
        [f"abstracts_{x}"], f"{s['speaker']}.")
    s2 = corpus.abstracts[y][1]
    add(f"Which track was the session \"{s2['title']}\" in at the {name[y]}?", "single_doc", ["session_abstracts"],
        [f"abstracts_{y}"], f"The {s2['track']} track.")
    add(f"Compare the operational issues at the {name[a]} and the {name[y]}.", "multi_doc", ["event_reports"],
        [f"report_{a}", f"report_{y}"], f"{name[a]}: {facts[a].issue}. {name[y]}: {facts[y].issue}.")
    r2 = results[results["event_id"] == y].sort_values("leads", ascending=False).iloc[0]
    price = next(p for p in _PACKAGE_INFO if p[0] == r2.package)
    year_y = int(pd.Timestamp(r2.start_date).year)
    add(f"Which package did {r2.sponsor} buy at the {name[y]}, and what does that package cost in {year_y}?",
        "multi_doc", ["event_reports", "pricing"], [f"report_{y}", f"pricing_{year_y}"],
        f"{r2.package}, which costs {_money(_year_price(price[2], year_y))} in {year_y}.")
    lesson_event = next(e for e in ids if facts[e].lesson == _LESSONS[0][0])
    add(f"What was the show-up rate at the {name[lesson_event]}, and what did the report recommend to improve it?",
        "mixed", ["event_reports"], [f"report_{lesson_event}"],
        f"{_pct(kpis.loc[lesson_event, 'show_up_rate'])}; the report recommends that we "
        f"{facts[lesson_event].recommendation}.")
    add(f"What was the hotel room rate for speakers at the {name[b]}?", "unanswerable", [], [],
        "Not in the documents.", answerable=False)
    add(f"Which catering company served the {name[c]}, and what did it charge?", "unanswerable", [], [],
        "Not in the documents.", answerable=False)
    add("What are the dates of next year's events?", "unanswerable", [], [], "Not in the documents.",
        answerable=False)
    return questions


# --- entry point ----------------------------------------------------------------------------------

def generate_documents(tables: dict[str, pd.DataFrame], data_dir: str | Path, seed: int = 42) -> list[DocRecord]:
    """Write the corpus, its manifest and the golden set under ``data_dir``; return the manifest records.

    ``tables`` must be the validated dataset (as ``core.schema.load_dataset`` returns it)."""
    data_dir = Path(data_dir)
    out = data_dir / DOCS_DIR
    for name in CORPORA:
        (out / name).mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng([seed, 2])  # independent of the CSV streams
    corpus = _Corpus()
    _event_reports(rng, tables, out, corpus)
    _pricing(out, corpus)
    _sponsor_decks(tables, out, corpus)
    _operations(out, corpus)
    _session_abstracts(rng, tables, out, corpus)

    (out / MANIFEST).write_text(json.dumps([asdict(r) for r in corpus.records], indent=2) + "\n",
                                encoding="utf-8", newline="\n")
    (data_dir / EVAL_DIR).mkdir(exist_ok=True)
    (data_dir / EVAL_DIR / GOLDEN_SET).write_text(json.dumps(_golden_set(tables, corpus), indent=2) + "\n",
                                                  encoding="utf-8", newline="\n")
    return corpus.records


def load_manifest(data_dir: str | Path) -> list[DocRecord]:
    path = Path(data_dir) / DOCS_DIR / MANIFEST
    return [DocRecord(**item) for item in json.loads(path.read_text(encoding="utf-8"))]

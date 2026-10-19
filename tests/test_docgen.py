"""The synthetic RAG corpus: every manifest file exists, figures match the analysis layer, the golden
set points at real documents, and the same seed gives the same content."""
import json
import unicodedata

import pytest

from core import analysis
from core.docgen import CORPORA, DOCS_DIR, EVAL_DIR, GOLDEN_SET, generate_documents, load_manifest


@pytest.fixture(scope="module")
def corpus(dataset, tmp_path_factory):
    out = tmp_path_factory.mktemp("corpus")
    generate_documents(dataset, out)
    return out


def _pdf_text(path) -> str:
    """Extracted text, NFKC-normalised: the PDF font uses ligatures, so "confirmed" extracts with a single
    "fi" character (U+FB01). Step 3's cleaner normalises the same way."""
    import pymupdf
    with pymupdf.open(path) as doc:
        return unicodedata.normalize("NFKC", "\n".join(page.get_text() for page in doc))


def test_manifest_lists_every_corpus_and_every_file_exists(corpus):
    records = load_manifest(corpus)
    assert {r.corpus for r in records} == set(CORPORA)
    assert len(records) == 12 + 2 + 6 + 4 + 12
    assert len({r.doc_id for r in records}) == len(records)
    for r in records:
        assert (corpus / DOCS_DIR / r.path).stat().st_size > 0, r.path


def test_report_figures_match_the_analysis_layer(corpus, dataset):
    k = analysis.event_kpis(dataset).set_index("event_id").loc["EV001"]
    text = " ".join(_pdf_text(corpus / DOCS_DIR / "event_reports" / "EV001_post_event_report.pdf").split())
    assert f"welcomed {k.attended:,} attendees out of {k.confirmed:,} confirmed registrations" in text
    results = analysis.sponsorship_results(dataset)
    for r in results[results["event_id"] == "EV001"].itertuples():
        assert r.sponsor in text
    assert "Page 1 of" in text  # the repeated footer that ingestion must strip


def test_pricing_2026_has_the_lead_guarantee_and_2025_does_not(corpus):
    p25 = _pdf_text(corpus / DOCS_DIR / "pricing" / "sponsorship_packages_2025.pdf")
    p26 = _pdf_text(corpus / DOCS_DIR / "pricing" / "sponsorship_packages_2026.pdf")
    assert "$48,000" in p25 and "$51,000" in p26 and "$800" in p26  # small add-ons round to $50
    assert "Lead guarantee" in p26 and "Lead guarantee" not in p25


def test_decks_and_sops_open_with_notes_and_numbered_steps(corpus):
    from docx import Document
    from pptx import Presentation
    deck = Presentation(str(next((corpus / DOCS_DIR / "sponsor_decks").glob("*.pptx"))))
    assert len(deck.slides) == 6
    assert "proposal" in deck.slides[5].notes_slide.notes_text_frame.text.lower()
    sop = Document(str(corpus / DOCS_DIR / "operations" / "booth_logistics.docx"))
    assert any("7:00 AM" in p.text and p.style.name == "List Number" for p in sop.paragraphs)


def test_golden_set_points_at_real_documents(corpus):
    questions = json.loads((corpus / EVAL_DIR / GOLDEN_SET).read_text(encoding="utf-8"))
    doc_ids = {r.doc_id for r in load_manifest(corpus)}
    assert len(questions) == 20 and len({q["id"] for q in questions}) == 20
    assert sum(not q["answerable"] for q in questions) == 3
    assert {q["category"] for q in questions} == {"single_doc", "table_lookup", "multi_doc", "mixed", "unanswerable"}
    for q in questions:
        assert set(q["expected_doc_ids"]) <= doc_ids, q["id"]
        assert q["answerable"] == bool(q["expected_doc_ids"]) and q["reference_answer"]


def test_same_seed_same_content(corpus, dataset, tmp_path):
    generate_documents(dataset, tmp_path)
    for rel in (f"{EVAL_DIR}/{GOLDEN_SET}", f"{DOCS_DIR}/manifest.json",
                f"{DOCS_DIR}/session_abstracts/EV005_session_abstracts.md"):
        assert (tmp_path / rel).read_text(encoding="utf-8") == (corpus / rel).read_text(encoding="utf-8")
    report = f"{DOCS_DIR}/event_reports/EV005_post_event_report.pdf"
    assert _pdf_text(tmp_path / report) == _pdf_text(corpus / report)


def test_bundled_corpus_is_up_to_date(corpus):
    """data/synthetic/docs must be regenerated (python scripts/generate_data.py) when docgen changes."""
    from core.config import DEFAULT_DATA_DIR
    bundled = json.loads((DEFAULT_DATA_DIR / EVAL_DIR / GOLDEN_SET).read_text(encoding="utf-8"))
    assert bundled == json.loads((corpus / EVAL_DIR / GOLDEN_SET).read_text(encoding="utf-8"))